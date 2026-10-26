"""Reglas de los ajustes solicitados, exclusivamente con datos ficticios."""
import sys,io,base64,unittest,json
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'APLICACION_ACTUAL'))
from test_phase2_domain import PhaseTwoTests
import app,phase2 as q,production as p,labflows as lab
class LabFlowsTests(PhaseTwoTests):
 def setUp(self):
  super().setUp();lab.install(app)
 def test_review_fields_and_required_sap(self):
  with app.database() as db:
   a=p._activity(db,'revision');keys={v['key'] for v in a['fields']}
   self.assertFalse(keys&{'modelo','tipo_inspeccion','valor_referencia','valor_medido','tolerancia_pct','fallo_detectado'})
   self.assertTrue(next(v for v in a['fields'] if v['key']=='sap')['required'])
   self.assertEqual(next(v for v in a['fields'] if v['key']=='resultado')['options'],['Reutilizable','No reutilizable','Pendiente'])
  with self.assertRaises(p.ApiError):self.event()
  self.event(data={'sap':'00123','clasificacion':'Nuevo','resultado':'Reutilizable'})
 def test_official_template_columns_title_and_dropdowns(self):
  import openpyxl
  with app.database() as db:
   for aid in lab.OFFICIAL:
    raw=q.template(db,self.admin,aid);w=openpyxl.load_workbook(io.BytesIO(raw));s=w['Captura']
    self.assertEqual(s['A1'].value,lab.TITLES[aid]);self.assertEqual([c.value for c in s[3]],[v for _,v in lab.OFFICIAL[aid]])
    self.assertTrue(s.data_validations.dataValidation)
 def test_sealing_is_bulk_only_and_assigns_row_employee(self):
  with self.assertRaises(p.ApiError):self.event(activity='sellado')
  row={'MEDIDOR':'000009','MARCA':'Marca ficticia','ARTICULO':'00123','MOVIMIENTO':'501','USUARIO':self.tech['name'],'FECHA':'2026-10-26','TIPO':'Nuevo','COLOR':'AZUL','SELLO #1':'001A','SELLO #2':'001B'}
  with app.database() as db:
   preview=q.handle_post(db,self.admin,'/api/p2/import/preview',{'activity_id':'sellado','rows':[row]});self.assertFalse(preview['errors'],preview)
   q.handle_post(db,self.admin,'/api/p2/import/confirm',{'preview_id':preview['preview_id']});r=db.execute('SELECT * FROM production_events').fetchone();self.assertEqual(r['responsible_id'],self.tech['id']);self.assertEqual(r['recorded_by'],self.admin['id'])
 def test_sap_is_movement_without_individual_ranking_and_file_duplicate(self):
  import openpyxl
  w=openpyxl.Workbook();s=w.active;s.append([lab.TITLES['sap_carga']]);s.append([v for _,v in lab.OFFICIAL['sap_carga']]);s.append(['00123','Material ficticio','3007','501',12,'2026-10-26']);b=io.BytesIO();w.save(b)
  data={'activity_id':'sap_carga','filename':'ficticio.xlsx','content':base64.b64encode(b.getvalue()).decode()}
  with app.database() as db:
   preview=q.handle_post(db,self.admin,'/api/p2/import/preview',data);self.assertFalse(preview['errors'],preview)
   q.handle_post(db,self.admin,'/api/p2/import/confirm',{'preview_id':preview['preview_id']})
   result=lab.analysis(db,self.admin,{'desde':'2026-10-26','hasta':'2026-10-26','activity_id':'sap_carga'});self.assertEqual(result['activities'][0]['quantity'],12);self.assertEqual(result['ranking'],[]);self.assertEqual(result['monthly'],[]);self.assertEqual(lab.report(result,'sap')[1][0][4],12)
   with self.assertRaises(p.ApiError):q.handle_post(db,self.admin,'/api/p2/import/preview',data)
 def test_retired_modules_block_writes_preserve_schema(self):
  with app.database() as db:
   self.assertFalse(p._activity(db,'bajas')['active']);self.assertFalse(p._activity(db,'asignacion')['active'])
   with self.assertRaises(p.ApiError):q.save_scrap(db,self.admin,{'serial':'00001','work_date':'2026-10-26','reason':'Prueba ficticia'})
 def test_review_preview_employee_and_reusable_statistics(self):
  row={'#':'1','FECHA':'2026-10-26','NUMERO DE EMPLEADO':self.tech['username'],'Medidor':'000002','MARCA':'Marca ficticia','SAP':'00123','VOLTAJE':'240V','CONDICIÓN':'Bueno','OBSERVACIONES':'Ficticio','resultado':'Reutilizable','clasificacion':'Nuevo'}
  with app.database() as db:
   preview=q.handle_post(db,self.admin,'/api/p2/import/preview',{'activity_id':'revision','rows':[row]});self.assertFalse(preview['errors'],preview);q.handle_post(db,self.admin,'/api/p2/import/confirm',{'preview_id':preview['preview_id']})
   result=lab.analysis(db,self.admin,{'desde':'2026-10-26','hasta':'2026-10-27'});self.assertEqual(result['ranking'][0]['responsible_id'],self.tech['id']);self.assertEqual(result['quality']['approval_pct'],100);self.assertEqual(result['matrix'][0]['average'],1);self.assertEqual(len(result['dates']),2)
 def test_official_review_file_and_source_preserved_on_revalidation(self):
  import openpyxl
  w=openpyxl.Workbook();sheet=w.active;sheet.append([lab.TITLES['revision']]);sheet.append([v for _,v in lab.OFFICIAL['revision']]);sheet.append(['1','2026-10-26',self.tech['username'],'000123','Marca ficticia','00123','240V','Bueno','Ficticio']);b=io.BytesIO();w.save(b)
  data={'activity_id':'revision','filename':'oficial.xlsx','content':base64.b64encode(b.getvalue()).decode(),'context':{'clasificacion':'Nuevo','resultado':'Reutilizable'}}
  with app.database() as db:
   preview=q.handle_post(db,self.admin,'/api/p2/import/preview',data);self.assertFalse(preview['errors'],preview)
   again=q.handle_post(db,self.admin,'/api/p2/import/preview',{'activity_id':'revision','rows':preview['source_rows'],'previous_preview_id':preview['preview_id']});self.assertFalse(again['errors'])
   q.handle_post(db,self.admin,'/api/p2/import/confirm',{'preview_id':again['preview_id']})
   with self.assertRaises(p.ApiError):q.handle_post(db,self.admin,'/api/p2/import/confirm',{'preview_id':preview['preview_id']})
   self.assertEqual(db.execute('SELECT count(*) FROM production_events').fetchone()[0],1)
 def test_statistics_tech_scope_previous_period_and_report_limits(self):
  self.event(day='2026-10-24',data={'sap':'00123','clasificacion':'Nuevo','resultado':'Reutilizable'})
  self.event(serial='00002',day='2026-10-26',data={'sap':'00123','clasificacion':'Nuevo','resultado':'No reutilizable'})
  with app.database() as db:
   result=lab.analysis(db,self.tech,{'desde':'2026-10-26','hasta':'2026-10-27'});self.assertEqual(result['activities'][0]['previous'],1);self.assertEqual(result['activities'][0]['growth_pct'],0);self.assertEqual(result['quality']['approval_pct'],0)
   other=lab.analysis(db,self.other,{'desde':'2026-10-26','hasta':'2026-10-27'});self.assertEqual(other['events'],0)
   big=lab.analysis(db,self.admin,{'desde':'2026-01-01','hasta':'2026-10-27'})
   with self.assertRaises(p.ApiError):lab.report(big,'daily')
 def test_historical_units_are_not_added_or_ranked_together(self):
  self.event(data={'sap':'00123','clasificacion':'Nuevo'})
  with app.database() as db, p._transaction(db):
   a=p._activity(db,'revision');p._save_activity(db,self.admin,{'id':'revision','name':a['name'],'mode':a['mode'],'unit':'otra unidad','fields':a['fields'],'active':True})
  self.event(serial='00009',day='2026-10-26',data={'sap':'00123','clasificacion':'Nuevo'})
  with app.database() as db:
   result=lab.analysis(db,self.admin,{'desde':'2026-10-24','hasta':'2026-10-26'});self.assertEqual(len(result['activities']),2);self.assertEqual([v['quantity'] for v in result['activities']],[1,1]);self.assertEqual([v['rank'] for v in result['ranking']],[1,1])
# Se hereda la preparación, no los casos de la especificación anterior.
for name in list(vars(PhaseTwoTests)):
 if name.startswith('test_') and name not in vars(LabFlowsTests):setattr(LabFlowsTests,name,None)
del PhaseTwoTests
if __name__=='__main__':unittest.main()
