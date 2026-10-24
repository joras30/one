"""Operación V34 con identidades sintéticas y bases efímeras, sin datos del laboratorio."""
import base64
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import uuid
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'APLICACION_ACTUAL'))
import app
import phase2 as q
import production as p

class PhaseTwoTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        app.configure(self.temp.name,production=True);app.init_db();self.key=app.foundation.bootstrap(app);q.install(app)
        with app.database() as db:
            self.admin=app.user_public(db,db.execute("SELECT * FROM users WHERE username='admin'").fetchone())
            for code,role in [('91001','tecnico'),('91002','tecnico'),('sup_uno','supervisor'),('sup_dos','supervisor')]:app.insert_user(db,code,'Persona ficticia '+code,role,'ClaveFicticia2026!')
            self.tech=app.user_public(db,db.execute("SELECT * FROM users WHERE username='91001'").fetchone())
            self.other=app.user_public(db,db.execute("SELECT * FROM users WHERE username='91002'").fetchone())
            self.supervisor=app.user_public(db,db.execute("SELECT * FROM users WHERE username='sup_uno'").fetchone())
            self.supervisor2=app.user_public(db,db.execute("SELECT * FROM users WHERE username='sup_dos'").fetchone())
            for kind,label in [('marcas','Marca ficticia'),('voltajes','240V'),('condiciones','Apto ficticio')]:q.catalog_save(db,self.admin,{'kind':kind,'label':label})
            app.foundation.put_setting(db,'catalog_ready',True)
            for val in q.VALUATIONS:q.save_article(db,self.admin,{'code':'00123','valuation':val,'description':'Material ficticio'})
            for activity in ('revision','sellado','bajas'):
                q.save_assignment(db,self.supervisor,{'user_id':self.tech['id'],'activity_id':activity,'permissions':['proponer_baja','ver'] if activity=='bajas' else ['ver','capturar','carga_masiva','panel'],'reason':'Asignación ficticia'})
    def event(self,serial='00001',activity='revision',day='2026-10-24',**kw):
        data={'serie':serial,'marca':'Marca ficticia','voltaje':'240V','condicion':'Apto ficticio'} if activity=='revision' else {'serie':serial,'marca':'Marca ficticia','sap':'00123','clasificacion':'Nuevo','sello_1':'000S1','sello_2':'000S2','color':'AZUL'}
        data.update(kw.pop('data',{}));payload={'activity_id':activity,'work_date':day,'quantity':1,'data':data,'operation_id':uuid.uuid4().hex,**kw}
        with app.database() as db:return p.handle_post(db,self.tech,'/api/eventos',payload)
    def test_reentries_are_silent_except_same_activity_condition_and_day(self):
        self.event();self.event(day='2026-10-25')
        with self.assertRaises(p.ApiError) as e:self.event(day='2026-10-25')
        self.assertTrue(e.exception.details['needs_duplicate_note'])
        self.event(day='2026-10-25',duplicate_note='Garantía ficticia')
        self.event(activity='sellado',day='2026-10-25')
        self.event(day='2026-10-26')
        with app.database() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM p2_cycles').fetchone()[0],2)
            self.assertEqual(db.execute('SELECT count(*) FROM production_events').fetchone()[0],5)
            self.assertTrue(app.production.operational.audit_integrity(db,self.admin)['ok'])
    def test_seals_are_unique_forever_and_same_number_other_color_is_allowed(self):
        self.event(activity='sellado')
        with self.assertRaises(p.ApiError) as e:self.event(serial='00002',activity='sellado')
        self.assertIn('seal_conflict',e.exception.details)
        with app.database() as db:
            old=db.execute('SELECT id FROM production_events').fetchone()[0]
            p._state_event(db,self.admin,old,{'reason':'Anulación ficticia'})
        with self.assertRaises(p.ApiError):self.event(activity='sellado')
        self.event(activity='sellado',data={'color':'MORADO'})
        with app.database() as db:self.assertEqual(db.execute('SELECT count(*) FROM p2_seals').fetchone()[0],4)
    def test_identical_pair_is_rejected_and_correction_cannot_reassign_a_seal(self):
        with self.assertRaises(p.ApiError):self.event(activity='sellado',data={'sello_2':'000S1'})
        result=self.event(activity='sellado')
        with app.database() as db:
            with self.assertRaises(p.ApiError):p._correct_event(db,self.admin,result['record']['id'],{'reason':'Prueba','data':{'sello_1':'nuevo'}})
    def test_every_supervisor_can_assign_every_technician_and_read_their_profile(self):
        with app.database() as db:
            app.foundation.assign_team(db,self.admin,self.other['id'],{'supervisor_id':self.supervisor2['id']})
            q.save_assignment(db,self.supervisor,{'user_id':self.other['id'],'activity_id':'revision','permissions':['capturar'],'reason':'Cobertura ficticia'})
            self.assertEqual(q.personnel(db,self.supervisor,self.other['id'])['person']['id'],self.other['id'])
            with self.assertRaises(p.ApiError):q.personnel(db,self.tech,self.other['id'])
            with self.assertRaises(p.ApiError):q.personnel(db,self.tech)
    def test_revocation_is_effective_and_does_not_grant_another_identity(self):
        with app.database() as db:
            db.execute('UPDATE p2_assignments SET revoked_at=?',(p.now(),))
        with self.assertRaises(p.ApiError):self.event()
        with app.database() as db:
            with self.assertRaises(p.ApiError):p._responsible(db,self.tech,{'responsible_id':self.other['id']},'revision')
    def test_scrap_requires_approval_keeps_history_blocks_operations_and_admin_can_revert(self):
        self.event()
        with app.database() as db:
            result=q.handle_post(db,self.tech,'/api/p2/scrap',{'serial':'00001','work_date':'2026-10-24','reason':'Irreparable ficticio','parts':[]})
            self.assertEqual(result['status'],'pendiente')
            self.assertEqual(db.execute("SELECT state FROM p2_meters").fetchone()[0],'activo')
            q.handle_post(db,self.supervisor2,'/api/p2/scrap/resolve',{'ids':[result['id']],'action':'autorizar'})
        with self.assertRaises(p.ApiError):self.event(day='2026-10-26')
        with app.database() as db:
            q.handle_post(db,self.admin,'/api/p2/scrap/resolve',{'ids':[result['id']],'action':'revertir','reason':'Revisión administrativa ficticia'})
            self.assertEqual(db.execute('SELECT count(*) FROM production_events').fetchone()[0],1)
        self.event(day='2026-10-26')
    def test_preview_detects_batch_seal_conflict_and_saves_nothing(self):
        rows=[{'MEDIDOR':str(i),'MARCA':'Marca ficticia','ARTICULO (SAP)':'00123','VALORACIÓN':'Nuevo','FECHA':'2026-10-24','COLOR':'AZUL','SELLO #1':'001','SELLO #2':str(10+i)} for i in (1,2)]
        with app.database() as db:
            result=q.handle_post(db,self.tech,'/api/p2/import/preview',{'activity_id':'sellado','rows':rows})
            self.assertEqual(len(result['errors']),1)
            self.assertEqual(db.execute('SELECT count(*) FROM production_events').fetchone()[0],0)
            self.assertEqual(db.execute('SELECT count(*) FROM p2_seals').fetchone()[0],0)
            with self.assertRaises(p.ApiError):q.handle_post(db,self.tech,'/api/p2/import/confirm',{'preview_id':result['preview_id']})
    def test_import_is_atomic_revalidated_and_idempotent(self):
        rows=[{'MEDIDOR':'00001','MARCA':'Marca ficticia','VOLTAJE':'240V','CONDICIÓN':'Apto ficticio','FECHA':'2026-10-24'}]
        with app.database() as db:
            result=q.handle_post(db,self.tech,'/api/p2/import/preview',{'activity_id':'revision','rows':rows});self.assertFalse(result['errors'])
            first=q.handle_post(db,self.tech,'/api/p2/import/confirm',{'preview_id':result['preview_id']})
            second=q.handle_post(db,self.tech,'/api/p2/import/confirm',{'preview_id':result['preview_id']})
            self.assertEqual(first['imported'],1);self.assertEqual(second['imported'],1)
            self.assertEqual(db.execute('SELECT count(*) FROM production_events').fetchone()[0],1)
    def test_template_is_real_excel_with_live_validations_reference_and_text_seals(self):
        import openpyxl
        with app.database() as db:raw=q.template(db,self.tech,'sellado')
        book=openpyxl.load_workbook(io.BytesIO(raw));self.assertEqual(book.sheetnames,['Captura','Guía de Referencia'])
        sheet=book['Captura'];self.assertGreater(len(sheet.data_validations.dataValidation),0)
        self.assertGreater(len(book.defined_names),0)
        index=[c.value for c in sheet[1]].index('SELLO #1')+1
        self.assertEqual(sheet.cell(2,index).number_format,'@')
    def test_sap_identity_is_pair_and_summary_has_exact_drilldown(self):
        self.event(activity='sellado')
        with app.database() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM p2_articles').fetchone()[0],2)
            report=q.sap_summary(db,self.supervisor,{})
            self.assertEqual((report['rows'][0]['quantity'],report['rows'][0]['warehouse']),(1,'3001'))
            self.assertEqual(report['rows'][0]['records'][0]['serial'],'00001')

    def test_import_revalidates_against_seals_used_after_preview_and_rolls_back_all_rows(self):
        rows=[{'MEDIDOR':'00001','MARCA':'Marca ficticia','ARTICULO (SAP)':'00123','VALORACIÓN':'Nuevo','FECHA':'2026-10-24','COLOR':'AZUL','SELLO #1':'nuevo1','SELLO #2':'nuevo2'},
              {'MEDIDOR':'00002','MARCA':'Marca ficticia','ARTICULO (SAP)':'00123','VALORACIÓN':'Nuevo','FECHA':'2026-10-24','COLOR':'AZUL','SELLO #1':'000S1','SELLO #2':'000S2'}]
        with app.database() as db:preview=q.handle_post(db,self.tech,'/api/p2/import/preview',{'activity_id':'sellado','rows':rows})
        self.assertFalse(preview['errors']);self.event(serial='00003',activity='sellado')
        with app.database() as db:
            with self.assertRaises(p.ApiError):q.handle_post(db,self.tech,'/api/p2/import/confirm',{'preview_id':preview['preview_id']})
            self.assertEqual(db.execute('SELECT count(*) FROM production_events').fetchone()[0],1)
            self.assertIsNone(db.execute("SELECT 1 FROM p2_seals WHERE number='nuevo1'").fetchone())

    def test_concurrent_seal_claims_produce_one_record_without_partial_writes(self):
        from concurrent.futures import ThreadPoolExecutor
        def worker(serial):
            try:self.event(serial=serial,activity='sellado');return True
            except p.ApiError:return False
        with ThreadPoolExecutor(max_workers=2) as executor:outcomes=list(executor.map(worker,['001','002']))
        self.assertEqual(sorted(outcomes),[False,True])
        with app.database() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM production_events').fetchone()[0],1)
            self.assertEqual(db.execute('SELECT count(*) FROM p2_seals').fetchone()[0],2)

    def test_bulk_permission_works_without_individual_capture_and_expiry_revokes_access(self):
        with app.database() as db:
            db.execute('UPDATE p2_assignments SET revoked_at=?',(p.now(),))
            q.save_assignment(db,self.supervisor,{'user_id':self.tech['id'],'activity_id':'revision','permissions':['carga_masiva'],'reason':'Solo carga masiva'})
            rows=[{'MEDIDOR':'00001','MARCA':'Marca ficticia','VOLTAJE':'240V','CONDICIÓN':'Apto ficticio','FECHA':'2026-10-24'}]
            preview=q.handle_post(db,self.tech,'/api/p2/import/preview',{'activity_id':'revision','rows':rows})
            self.assertFalse(preview['errors'])
            q.handle_post(db,self.tech,'/api/p2/import/confirm',{'preview_id':preview['preview_id']})
            db.execute('UPDATE p2_assignments SET ends_at=?',('2020-01-01T00:00:00Z',))
            with self.assertRaises(p.ApiError):q.template(db,self.tech,'revision')
        with self.assertRaises(p.ApiError):self.event(serial='00002')

    def test_excel_numeric_identifiers_are_rejected_instead_of_losing_zeroes(self):
        import openpyxl
        book=openpyxl.Workbook();sheet=book.active;sheet.append(['MEDIDOR','MARCA','VOLTAJE','CONDICIÓN','FECHA']);sheet.append([123,'Marca ficticia','240V','Apto ficticio','2026-10-24'])
        raw=io.BytesIO();book.save(raw)
        with app.database() as db:
            with self.assertRaises(ValueError):q.handle_post(db,self.tech,'/api/p2/import/preview',{'activity_id':'revision','filename':'prueba.xlsx','content':base64.b64encode(raw.getvalue()).decode()})
            self.assertEqual(db.execute('SELECT count(*) FROM production_events').fetchone()[0],0)

    def test_partial_receipts_keep_provider_balance_and_final_certificate(self):
        payload={'activity_id':'recepcion','work_date':'2026-10-24','quantity':3,'data':{'sap':'00123','clasificacion':'Nuevo','marca':'Marca ficticia','numero_orden':'ORDEN-FICTICIA','proceso':'PROCESO-FICTICIO','cantidad_esperada':5,'proveedor':'Proveedor ficticio'}}
        with app.database() as db:
            for quantity in (3,2):
                payload['operation_id']=uuid.uuid4().hex;payload['quantity']=quantity
                result=p.handle_post(db,self.admin,'/api/eventos',payload)
            order=p.operational.list_orders(db,self.admin,{})['orders'][0]
            self.assertEqual((order['received'],order['pending'],order['provider']),(5,0,'Proveedor ficticio'))
            result=p.operational.close_order(db,self.admin,order['id'],{})
            self.assertTrue(result['ok'])
            self.assertTrue(p.operational.certificate(db,self.admin,order['id']))

    def test_new_activity_has_versioned_catalog_fields_illustration_and_live_template(self):
        with app.database() as db:
            result=p.handle_post(db,self.admin,'/api/actividades',{'name':'Actividad ficticia nueva','mode':'individual','unit':'pruebas','illustration':'⚙','fields':[p._field('serie','Serie'),p._field('condicion','Condición','select',True,source='condiciones')]})
            activity=result['activity'];self.assertEqual(activity['illustration'],'⚙')
            q.save_assignment(db,self.supervisor,{'user_id':self.tech['id'],'activity_id':activity['id'],'permissions':['capturar','carga_masiva','panel'],'reason':'Prueba de actividad nueva'})
            self.assertTrue(q.template(db,self.tech,activity['id']).startswith(b'PK'))
            record=p.handle_post(db,self.tech,'/api/eventos',{'activity_id':activity['id'],'work_date':'2026-10-24','operation_id':uuid.uuid4().hex,'data':{'serie':'000NUEVO','condicion':'Apto ficticio'}})['record']
            p.handle_post(db,self.admin,'/api/actividades',{'id':activity['id'],'name':'Nombre nuevo ficticio','fields':activity['fields'],'illustration':'✓'})
            self.assertEqual(p._event(db,record['id'])['activity_name'],'Actividad ficticia nueva')

    def test_restoring_phase_one_backup_upgrades_without_replacing_master_credentials(self):
        self.event()
        with app.database() as db:
            # The first backup was taken before the phase-two activity migration.
            first=sorted((app.DATA_DIR/'backups').glob('*.db'))[0]
        preview=app.foundation.restore_preview(app,self.admin,{'path':str(first)})
        app.foundation.restore_confirm(app,self.admin,{'token':preview['token'],'confirm':'RESTAURAR','reason':'Restauración sintética de fase 1','password':self.key})
        with app.database() as db:
            self.assertTrue(q.enabled(db));self.assertEqual(app.foundation.setting(db,'validated_phase'),1)
            master=db.execute("SELECT * FROM users WHERE username='admin'").fetchone();self.assertTrue(app.verify_password(master,self.key))
            self.assertTrue(p.operational.audit_integrity(db,self.admin)['ok'])

    def test_scrap_retry_returns_existing_resolution_instead_of_proposing_twice(self):
        payload={'serial':'000B','work_date':'2026-10-24','reason':'Irreparable ficticio','parts':[],'operation_id':'baja-idempotente'}
        with app.database() as db:
            first=q.handle_post(db,self.tech,'/api/p2/scrap',payload)
            q.handle_post(db,self.supervisor,'/api/p2/scrap/resolve',{'ids':[first['id']],'action':'rechazar','reason':'Verificar condición'})
            again=q.handle_post(db,self.tech,'/api/p2/scrap',payload)
            self.assertEqual(again['status'],'rechazada');self.assertTrue(again['already_saved'])
            self.assertEqual(db.execute('SELECT count(*) FROM p2_scrap').fetchone()[0],1)

if __name__=='__main__':unittest.main()
